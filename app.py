import csv
import io
import os
import re
from io import BytesIO
from difflib import SequenceMatcher
from typing import Any

import openpyxl
import pyodbc
from flask import Flask, jsonify, redirect, render_template, request, url_for

SQL_SERVER_CONN_STR = os.getenv("PERSONNEL_SQL_CONNECTION_STRING", "").strip()

if not SQL_SERVER_CONN_STR:
    raise RuntimeError(
        "PERSONNEL_SQL_CONNECTION_STRING is required. SQLite fallback has been removed."
    )

app = Flask(__name__)

CSV_FIELD_ALIASES = {
    "employee_id": ["employee_id", "Employee ID", "employee id", "EmployeeId"],
    "name": ["name", "Name", "Worker", "worker", "Member"],
    "rank": ["rank", "Rank", "Worker/Position Job Profile", "Job Profile"],
    "division": ["division", "Division", "Assigned Division"],
    "email": ["email", "Email"],
    "status": ["status", "Status"],
    "badge_number": ["badge_number", "Badge Number", "badge number"],
    "sequence_num": ["Sequence", "Sequence Number", "sequence", "sequence number", "sequence_num"],
    "department_cell": ["Department Cell", "department cell", "department_cell"],
    "radio_id": ["Radio_ID", "Radio ID", "radio_id", "radio id"],
    "race": ["Race", "race"],
    "sex": ["Sex", "sex"],
}

PENDING_APPROVALS: dict[str, dict[str, str]] = {}




DIVISION_GROUPS = {
    "Command": ["Command", "Command Staff"],
    "Court Security": ["Court Security", "Court_Security"],
    "Domestic Violence": ["Domestic Violence", "Domestic_Violence"],
    "Special Operations": ["Special Operations", "Special_Operations", "SOD"],
}

DIVISION_OPTIONS = [
    "Command",
    "Communications",
    "Court Security",
    "Domestic Violence",
    "Field Services",
    "Human Resources",
    "Information Technology",
    "Internal Affairs",
    "Neighborhood Services",
    "Special Operations",
    "Training",
    "Others",
]
STATUS_OPTIONS = ["Active", "Inactive"]
RACE_OPTIONS = ["White", "Black", "Asian", "Hispanic", "Other"]
SEX_OPTIONS = ["Male", "Female", "Other"]


def normalize_division_value(value: str) -> str:
    return (value or "").replace("_", " ").strip()


def display_division(value: str) -> str:
    normalized = normalize_division_value(value)
    for label, variants in DIVISION_GROUPS.items():
        normalized_variants = {normalize_division_value(v).lower() for v in variants}
        if normalized.lower() in normalized_variants:
            return label
    return normalized


EDITABLE_FIELDS = {
    "name",
    "email",
    "rank",
    "division",
    "status",
    "sequence_num",
    "department_cell",
    "radio_id",
    "race",
    "sex",
    "notes",
}


def normalize_name(name: str) -> str:
    cleaned = re.sub(r"[^a-z0-9\s]", " ", (name or "").lower())
    tokens = [token for token in cleaned.split() if len(token) > 1]
    return " ".join(tokens)


def swapped_name_variants(name: str) -> set[str]:
    normalized = normalize_name(name)
    if not normalized:
        return set()

    parts = normalized.split()
    variants = {normalized}
    if len(parts) >= 2:
        variants.add(" ".join([parts[-1], *parts[:-1]]))
    return {v for v in variants if v}


def build_name_index(db: Any) -> dict[str, tuple[str, str]]:
    cursor = db.cursor()
    cursor.execute("SELECT employee_id, name FROM dbo.agency_members WHERE name IS NOT NULL")
    name_index: dict[str, tuple[str, str]] = {}
    for emp_id, name in cursor.fetchall():
        existing_name = str(name or "")
        for v in swapped_name_variants(existing_name):
            if v not in name_index:
                name_index[v] = (str(emp_id), existing_name)
    return name_index


def best_fuzzy_match(name: str, name_index: dict[str, tuple[str, str]]) -> tuple[str, str, str, float]:
    candidates = swapped_name_variants(name)
    best_emp_id = ""
    best_variant = ""
    best_existing_name = ""
    best_score = 0.0

    for candidate in candidates:
        for existing_variant, (emp_id, existing_name) in name_index.items():
            score = SequenceMatcher(None, candidate, existing_variant).ratio()
            if score > best_score:
                best_score = score
                best_variant = existing_variant
                best_emp_id = emp_id
                best_existing_name = existing_name

    return best_emp_id, best_existing_name, best_variant, best_score


def get_db() -> Any:
    return pyodbc.connect(SQL_SERVER_CONN_STR)


def initialize_database(db: Any) -> None:
    cursor = db.cursor()
    cursor.execute(
        """
        IF OBJECT_ID('dbo.agency_members', 'U') IS NULL
        BEGIN
            CREATE TABLE dbo.agency_members (
                employee_id NVARCHAR(50) NOT NULL PRIMARY KEY,
                name NVARCHAR(200) NULL,
                email NVARCHAR(320) NULL,
                rank NVARCHAR(200) NULL,
                division NVARCHAR(200) NULL,
                status NVARCHAR(100) NULL,
                badge_number NVARCHAR(100) NULL,
                source_file NVARCHAR(260) NULL,
                imported_at DATETIME2 NOT NULL DEFAULT SYSUTCDATETIME()
            );
        END;
        """
    )
    for stmt in [
        "IF COL_LENGTH('dbo.agency_members','source_file') IS NULL ALTER TABLE dbo.agency_members ADD source_file NVARCHAR(260) NULL;",
        "IF COL_LENGTH('dbo.agency_members','imported_at') IS NULL ALTER TABLE dbo.agency_members ADD imported_at DATETIME2 NOT NULL DEFAULT SYSUTCDATETIME();",
        "IF COL_LENGTH('dbo.agency_members','sequence_num') IS NULL ALTER TABLE dbo.agency_members ADD sequence_num NVARCHAR(50) NULL;",
        "IF COL_LENGTH('dbo.agency_members','department_cell') IS NULL ALTER TABLE dbo.agency_members ADD department_cell NVARCHAR(100) NULL;",
        "IF COL_LENGTH('dbo.agency_members','radio_id') IS NULL ALTER TABLE dbo.agency_members ADD radio_id NVARCHAR(50) NULL;",
        "IF COL_LENGTH('dbo.agency_members','race') IS NULL ALTER TABLE dbo.agency_members ADD race NVARCHAR(50) NULL;",
        "IF COL_LENGTH('dbo.agency_members','sex') IS NULL ALTER TABLE dbo.agency_members ADD sex NVARCHAR(20) NULL;",
        "IF COL_LENGTH('dbo.agency_members','notes') IS NULL ALTER TABLE dbo.agency_members ADD notes NVARCHAR(MAX) NULL;",
    ]:
        cursor.execute(stmt)
    db.commit()


def _get_field(row: dict[str, str], canonical_name: str) -> str:
    for key in CSV_FIELD_ALIASES[canonical_name]:
        value = row.get(key)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def _update_member(cursor: Any, payload: dict[str, str], employee_id: str) -> None:
    cursor.execute(
        """
        UPDATE dbo.agency_members
        SET name = ?, email = ?, rank = ?, division = ?, status = ?, badge_number = ?,
            sequence_num = ?, department_cell = ?, radio_id = ?, race = ?, sex = ?,
            source_file = ?, imported_at = SYSUTCDATETIME()
        WHERE employee_id = ?
        """,
        payload["name"], payload["email"], payload["rank"], payload["division"], payload["status"],
        payload["badge_number"], payload["sequence_num"], payload["department_cell"], payload["radio_id"],
        payload["race"], payload["sex"], payload["source_file"], employee_id,
    )


def ingest_csv_stream(db: Any, csv_stream: io.TextIOBase, source_name: str) -> tuple[int, int, int, list[str]]:
    inserted = 0
    updated = 0
    skipped = 0
    logs: list[str] = []
    reader = csv.DictReader(csv_stream)
    cursor = db.cursor()
    name_index = build_name_index(db)

    for i, row in enumerate(reader, start=1):
        payload = {k: _get_field(row, k) for k in CSV_FIELD_ALIASES.keys()}
        payload["source_file"] = source_name
        target_employee_id = payload["employee_id"]

        if target_employee_id:
            cursor.execute("SELECT employee_id FROM dbo.agency_members WHERE employee_id = ?", target_employee_id)
            existing = cursor.fetchone()
            if existing:
                logs.append(f"MATCH employee_id: {target_employee_id} -> update")
            else:
                logs.append(f"NO employee_id match: {target_employee_id}")
        else:
            existing = None

        if not existing and payload["name"]:
            for variant in swapped_name_variants(payload["name"]):
                if variant in name_index:
                    target_employee_id = name_index[variant][0]
                    existing = (target_employee_id,)
                    logs.append(f"MATCH name: '{payload['name']}' -> employee_id {target_employee_id}")
                    break

        if not existing and payload["name"]:
            fuzzy_emp_id, fuzzy_existing_name, fuzzy_variant, fuzzy_score = best_fuzzy_match(payload["name"], name_index)
            if fuzzy_score >= 0.8 and fuzzy_emp_id:
                approval_id = f"{source_name}:{i}:{payload['name']}"
                PENDING_APPROVALS[approval_id] = {
                    **payload,
                    "employee_id": fuzzy_emp_id,
                    "name": payload["name"],
                    "score": f"{fuzzy_score:.2f}",
                    "variant": fuzzy_variant,
                    "suggested_name": fuzzy_existing_name,
                }
                logs.append(
                    f"NO exact match: '{payload['name']}' | suggestion employee_id={fuzzy_emp_id} name='{fuzzy_existing_name}' score={fuzzy_score:.2f}"
                )
                continue

        if existing:
            _update_member(cursor, payload, target_employee_id)
            updated += 1
            for v in swapped_name_variants(payload["name"]):
                name_index[v] = (target_employee_id, payload["name"])
        else:
            if not payload["employee_id"]:
                skipped += 1
                logs.append(f"SKIP row {i}: no employee_id and no qualified fuzzy match")
                continue
            cursor.execute(
                """
                INSERT INTO dbo.agency_members
                (employee_id, name, email, rank, division, status, badge_number, sequence_num,
                 department_cell, radio_id, race, sex, source_file, imported_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, SYSUTCDATETIME())
                """,
                payload["employee_id"], payload["name"], payload["email"], payload["rank"], payload["division"],
                payload["status"], payload["badge_number"], payload["sequence_num"], payload["department_cell"],
                payload["radio_id"], payload["race"], payload["sex"], payload["source_file"],
            )
            inserted += 1
            logs.append(f"INSERT employee_id: {payload['employee_id']}")

    db.commit()
    no_match_logs = [line for line in logs if line.startswith("NO ") or line.startswith("SKIP ")]
    matched_logs = [line for line in logs if line not in no_match_logs]
    return inserted, updated, skipped, no_match_logs + matched_logs + [f"SUMMARY inserted={inserted} updated={updated} skipped={skipped}"]


def fetch_members(db: Any, search_name: str = "", search_division: str = "", search_radio_id: str = ""):
    cursor = db.cursor()
    where_clauses = []
    params: list[str] = []

    if search_name:
        where_clauses.append("name LIKE ?")
        params.append(f"%{search_name}%")
    if search_division:
        if search_division in DIVISION_GROUPS:
            variants = DIVISION_GROUPS[search_division]
            placeholders = ",".join(["?"] * len(variants))
            where_clauses.append(f"REPLACE(division, '_', ' ') IN ({placeholders})")
            params.extend([normalize_division_value(v) for v in variants])
        else:
            where_clauses.append("REPLACE(division, '_', ' ') = ?")
            params.append(normalize_division_value(search_division))
    if search_radio_id:
        where_clauses.append("radio_id LIKE ?")
        params.append(f"%{search_radio_id}%")

    where_sql = f" WHERE {' AND '.join(where_clauses)}" if where_clauses else ""
    query = "SELECT employee_id, name, email, rank, division, status, badge_number, sequence_num, department_cell, radio_id, race, sex, notes FROM dbo.agency_members" + where_sql + " ORDER BY name, employee_id"
    cursor.execute(query, params)
    rows = cursor.fetchall()
    return [{"employee_id": r[0], "name": r[1], "email": r[2], "rank": r[3], "division": r[4], "division_display": display_division(r[4]), "status": r[5], "badge_number": r[6], "sequence_num": r[7], "department_cell": r[8], "radio_id": r[9], "race": r[10], "sex": r[11], "notes": r[12]} for r in rows]




def extract_email_rows_from_workbook(file_bytes: bytes) -> list[dict[str, str]]:
    wb = openpyxl.load_workbook(filename=BytesIO(file_bytes), data_only=True)
    rows: list[dict[str, str]] = []

    for ws in wb.worksheets:
        division = (ws.title or "").strip()
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0

        for r in range(1, max_row + 1):
            values = []
            for c in range(1, max_col + 1):
                v = ws.cell(row=r, column=c).value
                values.append(str(v).strip() if v is not None else "")

            emails = [v for v in values if "@" in v and "." in v]
            if not emails:
                continue

            email = emails[0]
            email_idx = values.index(email)
            candidate_name = ""
            for i in range(email_idx - 1, -1, -1):
                txt = values[i]
                if txt and txt.lower() not in {"rank", "deputies", "email", "name"}:
                    candidate_name = txt
                    break

            if not candidate_name:
                continue

            rows.append({"name": candidate_name, "email": email, "division": division})

    return rows


def ingest_email_workbook(db: Any, file_bytes: bytes, source_name: str) -> tuple[int, int, int, list[str]]:
    updates = 0
    approvals = 0
    skipped = 0
    logs: list[str] = []

    cursor = db.cursor()
    name_index = build_name_index(db)
    workbook_rows = extract_email_rows_from_workbook(file_bytes)

    for i, item in enumerate(workbook_rows, start=1):
        name = item["name"]
        email = item["email"]
        division = item["division"]
        target_employee_id = ""
        existing = None

        for variant in swapped_name_variants(name):
            if variant in name_index:
                target_employee_id = name_index[variant][0]
                existing = (target_employee_id,)
                logs.append(f"EMAIL MATCH name: '{name}' -> employee_id {target_employee_id}")
                break

        if not existing:
            fuzzy_emp_id, fuzzy_existing_name, fuzzy_variant, fuzzy_score = best_fuzzy_match(name, name_index)
            if fuzzy_score >= 0.8 and fuzzy_emp_id:
                approval_id = f"emailwb:{source_name}:{i}:{name}"
                PENDING_APPROVALS[approval_id] = {
                    "employee_id": fuzzy_emp_id,
                    "name": name,
                    "email": email,
                    "division": division,
                    "rank": "",
                    "status": "",
                    "sequence_num": "",
                    "department_cell": "",
                    "radio_id": "",
                    "race": "",
                    "sex": "",
                    "badge_number": "",
                    "source_file": source_name,
                    "score": f"{fuzzy_score:.2f}",
                    "variant": fuzzy_variant,
                    "suggested_name": fuzzy_existing_name,
                }
                approvals += 1
                logs.append(f"EMAIL NO exact match: '{name}' | suggestion '{fuzzy_existing_name}' score={fuzzy_score:.2f}")
            else:
                skipped += 1
                logs.append(f"EMAIL SKIP: '{name}' no match")
            continue

        cursor.execute(
            """
            UPDATE dbo.agency_members
            SET email = ?, division = ?, source_file = ?, imported_at = SYSUTCDATETIME()
            WHERE employee_id = ?
            """,
            email,
            division,
            source_name,
            target_employee_id,
        )
        updates += 1

    db.commit()
    logs.append(f"EMAIL SUMMARY updated={updates} approvals={approvals} skipped={skipped}")
    return updates, approvals, skipped, logs


def fetch_name_suggestions(db: Any, q: str, limit: int = 10) -> list[str]:
    q = q.strip()
    if len(q) < 3:
        return []
    cursor = db.cursor()
    cursor.execute(
        """
        SELECT DISTINCT TOP (?) name
        FROM dbo.agency_members
        WHERE name IS NOT NULL AND name LIKE ?
        ORDER BY name
        """,
        limit,
        f"%{q}%",
    )
    return [row[0] for row in cursor.fetchall() if row[0]]



def fetch_rank_options(db: Any) -> list[str]:
    cursor = db.cursor()
    cursor.execute(
        """
        SELECT DISTINCT rank
        FROM dbo.agency_members
        WHERE rank IS NOT NULL AND LTRIM(RTRIM(rank)) <> ''
        ORDER BY rank
        """
    )
    ranks = [str(row[0]).strip() for row in cursor.fetchall() if str(row[0]).strip()]
    return [rank for rank in ranks if rank != "Others"] + ["Others"]

def fetch_divisions(db: Any) -> list[str]:
    cursor = db.cursor()
    cursor.execute("SELECT DISTINCT division FROM dbo.agency_members WHERE division IS NOT NULL AND LTRIM(RTRIM(division)) <> '' ORDER BY division")
    raw = [row[0] for row in cursor.fetchall()]
    mapped = {display_division(v) for v in raw if v}
    return sorted(mapped)






@app.get("/name-suggestions")
def name_suggestions():
    q = request.args.get("q", "")
    try:
        db = get_db()
        names = fetch_name_suggestions(db, q)
        db.close()
        return jsonify({"suggestions": names})
    except Exception as exc:
        return jsonify({"suggestions": [], "error": str(exc)}), 500

@app.post("/update-cell")
def update_cell():
    payload = request.get_json(silent=True) or {}
    employee_id = str(payload.get("employee_id", "")).strip()
    field = str(payload.get("field", "")).strip()
    value = str(payload.get("value", "")).strip()

    if not employee_id or field not in EDITABLE_FIELDS:
        return jsonify({"ok": False, "error": "Invalid employee_id or field."}), 400

    db = get_db()
    cursor = db.cursor()
    cursor.execute(
        f"UPDATE dbo.agency_members SET {field} = ?, imported_at = SYSUTCDATETIME() WHERE employee_id = ?",
        value,
        employee_id,
    )
    db.commit()
    db.close()
    return jsonify({"ok": True})

@app.post("/approve")
def approve_guess():
    approval_id = request.form.get("approval_id", "")
    if approval_id in PENDING_APPROVALS:
        payload = PENDING_APPROVALS.pop(approval_id)
        db = get_db()
        cursor = db.cursor()
        _update_member(cursor, payload, payload["employee_id"])
        db.commit()
        db.close()
    return redirect(url_for("index"))


@app.post("/add-member")
def add_member():
    employee_id = request.form.get("employee_id", "").strip()
    if not employee_id:
        return redirect(url_for("index"))

    fields = {
        "name": request.form.get("name", "").strip(),
        "email": request.form.get("email", "").strip(),
        "rank": request.form.get("rank", "").strip(),
        "division": request.form.get("division", "").strip(),
        "status": request.form.get("status", "").strip(),
        "sequence_num": request.form.get("sequence_num", "").strip(),
        "department_cell": request.form.get("department_cell", "").strip(),
        "radio_id": request.form.get("radio_id", "").strip(),
        "race": request.form.get("race", "").strip(),
        "sex": request.form.get("sex", "").strip(),
        "notes": request.form.get("notes", "").strip(),
    }

    db = get_db()
    initialize_database(db)
    cursor = db.cursor()
    cursor.execute(
        """
        INSERT INTO dbo.agency_members
        (employee_id, name, email, rank, division, status, sequence_num, department_cell,
         radio_id, race, sex, notes, source_file, imported_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, SYSUTCDATETIME())
        """,
        employee_id,
        fields["name"],
        fields["email"],
        fields["rank"],
        fields["division"],
        fields["status"],
        fields["sequence_num"],
        fields["department_cell"],
        fields["radio_id"],
        fields["race"],
        fields["sex"],
        fields["notes"],
        "manual",
    )
    db.commit()
    db.close()
    return redirect(url_for("index"))


@app.route("/", methods=["GET", "POST"])
def index():
    import_result = None
    db_error = None
    members = []
    divisions: list[str] = []
    rank_options: list[str] = ["Others"]
    ingest_logs: list[str] = []

    search_name = request.values.get("search_name", "").strip()
    search_division = request.values.get("search_division", "").strip()
    search_radio_id = request.values.get("search_radio_id", "").strip()
    try:
        db = get_db()
        initialize_database(db)
        if request.method == "POST":
            uploaded_email_wb = request.files.get("email_workbook")
            uploaded = request.files.get("csv_file")
            if uploaded_email_wb and uploaded_email_wb.filename:
                wb_bytes = uploaded_email_wb.read()
                updates, approvals, skipped, ingest_logs = ingest_email_workbook(db, wb_bytes, uploaded_email_wb.filename)
                import_result = f"Email workbook import complete. Updated: {updates}, Pending approvals: {approvals}, Skipped: {skipped}."
            elif uploaded and uploaded.filename:
                text_stream = io.TextIOWrapper(uploaded.stream, encoding="utf-8-sig", newline="")
                inserted, updated, skipped, ingest_logs = ingest_csv_stream(db, text_stream, uploaded.filename)
                import_result = f"Import complete to Azure SQL. Inserted: {inserted}, Updated: {updated}, Skipped: {skipped}."
            else:
                import_result = "Please choose a CSV file or Email Workbook before clicking upload."
        members = fetch_members(db, search_name=search_name, search_division=search_division, search_radio_id=search_radio_id)
        divisions = fetch_divisions(db)
        rank_options = fetch_rank_options(db)
        db.close()
    except Exception as exc:
        db_error = str(exc)

    return render_template("index.html", members=members, import_result=import_result, db_error=db_error, ingest_logs=ingest_logs, pending_approvals=PENDING_APPROVALS, divisions=divisions, division_options=DIVISION_OPTIONS, status_options=STATUS_OPTIONS, rank_options=rank_options, race_options=RACE_OPTIONS, sex_options=SEX_OPTIONS, search_name=search_name, search_division=search_division, search_radio_id=search_radio_id)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8000, debug=True)
